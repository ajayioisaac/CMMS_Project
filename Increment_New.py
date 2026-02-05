from openpyxl import load_workbook
from datetime import timedelta, datetime, date
from dateutil.relativedelta import relativedelta

file_path = r"D:/CESL/CMMS Issues/Refined/Preventive_Mainenance_data_2026-01-31T08_29_13.606Z.xlsx"


def swap(sheet):
    for row in range(2, sheet.max_row + 1):
        gp_value = sheet[f"BW{row}"].value

        if isinstance(gp_value, (datetime, date)):
            sheet[f"H{row}"] = gp_value


def to_date(value):
    if isinstance(value, (datetime, date)):
        return value
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%m/%d/%Y"):
            try:
                return datetime.strptime(value, fmt)
            except ValueError:
                pass
    return None


def increment(sheet):
    for row in range(2, sheet.max_row + 1):
        last_done_raw = sheet[f"H{row}"].value
        occurrence = sheet[f"G{row}"].value
        frequency = sheet[f"F{row}"].value

        last_done = to_date(last_done_raw)

        if not last_done or not occurrence or not frequency:
            continue

        occurrence = occurrence.strip().lower()

        if occurrence == "week":
            next_due = last_done + timedelta(weeks=int(frequency))
        elif occurrence == "month":
            next_due = last_done + relativedelta(months=int(frequency))
        elif occurrence == "day":
            next_due = last_done + timedelta(days=int(frequency))
        elif occurrence == "year":
            next_due = last_done + relativedelta(years=int(frequency))
        else:
            continue

        sheet[f"I{row}"] = next_due

def clean_swap():
    wb = load_workbook(file_path)
    sheet = wb["Preventive_Mainenance_data_2026"]

    swap(sheet)
    increment(sheet)

    wb.save(file_path)
    wb.close()
    print("Swap and increment completed successfully.")

if __name__ == "__main__":
    clean_swap()
